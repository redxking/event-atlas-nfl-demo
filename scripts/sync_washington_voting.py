"""Import Washington SOS November 2026 drop boxes and voting centers."""
import hashlib
import io
import json
import re
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SNAPSHOT = ROOT / 'data/voting_locations.json'
SOURCE_ID = 'wa-nov-2026-sites'
PAGE = 'https://www.sos.wa.gov/elections/voters/voter-registration/drop-box-and-voting-center-locations'
XLSX = 'https://www.sos.wa.gov/sites/default/files/2026-01/Drop%20Boxes_0.xlsx'
NOTE = ('Washington SOS statewide November 2026 General Election site list. Workbook has site addresses and points but no per-site dates or hours. '
        'County submissions can change; verify current operation and schedule with the county elections office.')


def clean(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip()


def fetch():
    request = urllib.request.Request(XLSX, headers={'User-Agent': 'Mozilla/5.0 EventAtlas public-data evaluation'})
    with urllib.request.urlopen(request, timeout=30) as response:
        return response.read()


def parse(data, retrieved):
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if workbook.sheetnames != ['Sheet1']:
        raise ValueError(f'Unexpected Washington workbook sheets: {workbook.sheetnames}')
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    if tuple(header[:8]) != ('County', 'Place', 'Type', 'Address', 'Description', 'ppLatitude', 'ppLongitude', 'Zip Code'):
        raise ValueError(f'Unexpected Washington workbook header: {header[:8]}')
    edited = workbook.properties.modified
    edited_at = edited.isoformat() + 'Z' if edited else None
    result = []
    counties = set()
    for line, cells in enumerate(rows, 2):
        county, name, kind, address, description, lat, lon, postal = cells[:8]
        county, name, kind, address = (clean(v) for v in (county, name, kind, address))
        if not all((county, name, kind, address)) or kind not in ('Drop Box', 'Voting Center'):
            raise ValueError(f'Incomplete or unknown Washington site on workbook row {line}')
        if not isinstance(lat, (int, float)) or not isinstance(lon, (int, float)) or not (45 <= lat <= 50 and -125 <= lon <= -116):
            raise ValueError(f'Invalid Washington point on workbook row {line}')
        counties.add(county)
        digest = hashlib.sha256('|'.join((county, name, kind, address)).encode()).hexdigest()[:20]
        result.append(dict(id=f'{SOURCE_ID}:{digest}', sourceId=SOURCE_ID, sourceRecordId=digest,
                           name=name, type='ballot_drop_box' if kind == 'Drop Box' else 'voting_center',
                           jurisdiction='WA', county=county, city='', street=address, addressLine2='',
                           postal=clean(postal), lat=lat, lon=lon, precinct=None, ward='',
                           status='Published for November 3, 2026 General Election', datesOpen='', hours='',
                           votingSpace='', accessibility='', description=clean(description),
                           sourceUrl=PAGE, sourceRecordUrl=XLSX, sourceDataStatus=NOTE,
                           retrievedAt=retrieved, sourceEditedAt=edited_at,
                           coordinateStatus='Publisher workbook point', sourceRow=line))
    if len(result) < 600 or len(counties) != 39 or len({row['id'] for row in result}) != len(result):
        raise ValueError(f'Washington workbook layout/count changed: {len(result)} rows, {len(counties)} counties')
    return result, edited_at


def main():
    snapshot = json.loads(SNAPSHOT.read_text())
    retrieved = datetime.now(timezone.utc).isoformat().replace('+00:00', 'Z')
    try:
        rows, edited_at = parse(fetch(), retrieved)
        source = dict(id=SOURCE_ID, name='Washington November 2026 drop boxes and voting centers',
                      dataset=XLSX, publisherPage=PAGE, records=len(rows), reportedCount=len(rows),
                      retrievedAt=retrieved, dataEditedAt=edited_at, coverageNote=NOTE, status='ok')
        print(f'{SOURCE_ID}: {len(rows)} rows across 39 counties')
    except Exception as error:
        rows = [row for row in snapshot['locations'] if row['sourceId'] == SOURCE_ID]
        previous = next((source for source in snapshot['sources'] if source['id'] == SOURCE_ID), None)
        source = dict(id=SOURCE_ID, name='Washington November 2026 drop boxes and voting centers',
                      dataset=XLSX, publisherPage=PAGE, records=len(rows), retrievedAt=retrieved,
                      lastSuccessfulAt=(previous or {}).get('retrievedAt') if rows else None,
                      coverageNote=NOTE, status='stale_retained' if rows else 'error', error=str(error))
        print(f'{SOURCE_ID}: {error}; retained {len(rows)} prior rows')
    snapshot['sources'] = [source for source in snapshot['sources'] if source['id'] != SOURCE_ID] + [source]
    snapshot['locations'] = [row for row in snapshot['locations'] if row['sourceId'] != SOURCE_ID] + rows
    snapshot['retrievedAt'] = retrieved
    snapshot['coverageNote'] = ('Official-source site-type records for NY, DC, MD, PA, NC and WA. Washington November 2026 lists statewide '
                                'drop boxes and voting centers without per-site schedules. Other coverage is dated or incomplete. '
                                'Verify current site status with election officials.')
    temp = SNAPSHOT.with_suffix('.json.tmp')
    temp.write_text(json.dumps(snapshot, ensure_ascii=False, separators=(',', ':')))
    temp.replace(SNAPSHOT)
    print(f'Wrote {len(snapshot["locations"])} voting-site rows')


if __name__ == '__main__':
    main()
